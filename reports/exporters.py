"""
Both exporters take the exact same input shape every function in
services.py returns — {"title", "columns", "rows"} — so adding a seventh
report type later never requires touching this file.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def export_excel(report):
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:31] or "Report"  # Excel sheet-name limit

    ws.append([report["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    header_row_idx = 3
    ws.append(report["columns"])
    header_fill = PatternFill(start_color="2E86DE", end_color="2E86DE", fill_type="solid")
    for cell in ws[header_row_idx]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in report["rows"]:
        ws.append(row)

    for column_cells in ws.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        ws.column_dimensions[column_cells[0].column_letter].width = min(40, max(12, length + 2))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_pdf(report):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    elements = [Paragraph(report["title"], styles["Title"]), Spacer(1, 16)]

    table_data = [report["columns"]] + [[str(cell) for cell in row] for row in report["rows"]]
    if len(table_data) == 1:
        table_data.append(["No data for this period"] + [""] * (len(report["columns"]) - 1))

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E86DE")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer
